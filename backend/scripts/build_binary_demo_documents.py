#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Генерирует DOCX/PDF/XLSX/PPTX входные документы для release go/no-go demo")
    parser.add_argument(
        "--output-dir",
        default="backend/examples/cases/release_go_no_go_multifile_case/input",
        help="Директория input demo-кейса",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Перезаписать существующие binary demo files",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    written = build_binary_demo_documents(
        output_dir=Path(args.output_dir),
        overwrite=args.overwrite,
    )

    for path in written:
        print(f"written: {path}")
    if not written:
        print("binary_demo_documents: already_exist")


def build_binary_demo_documents(*, output_dir: Path, overwrite: bool = False) -> list[str]:
    output_dir.mkdir(parents=True, exist_ok=True)

    docx_path = output_dir / "05_release_notes.docx"
    pdf_path = output_dir / "06_audit_summary.pdf"
    scanned_pdf_path = output_dir / "07_scanned_signoff.pdf"
    scanned_ocr_sidecar_path = output_dir / "07_scanned_signoff.pdf.ocr.txt"
    xlsx_path = output_dir / "08_release_tracker.xlsx"
    pptx_path = output_dir / "09_release_briefing.pptx"

    written: list[str] = []
    if overwrite or not docx_path.exists():
        _write_docx(docx_path)
        written.append(str(docx_path))
    if overwrite or not pdf_path.exists():
        _write_pdf(pdf_path)
        written.append(str(pdf_path))
    if overwrite or not scanned_pdf_path.exists():
        _write_scanned_pdf(scanned_pdf_path)
        written.append(str(scanned_pdf_path))
    if overwrite or not scanned_ocr_sidecar_path.exists():
        _write_scanned_pdf_sidecar(scanned_ocr_sidecar_path)
        written.append(str(scanned_ocr_sidecar_path))
    if overwrite or not xlsx_path.exists():
        _write_xlsx(xlsx_path)
        written.append(str(xlsx_path))
    if overwrite or not pptx_path.exists():
        _write_pptx(pptx_path)
        written.append(str(pptx_path))
    return written


def _write_docx(path: Path) -> None:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("Для генерации DOCX demo требуется зависимость python-docx") from exc

    document = Document()
    document.add_heading("Release Notes: Payments v2", level=1)
    document.add_heading("Deployment Scope", level=2)
    document.add_paragraph("Payments v2 release includes API rollout, monitoring updates, and support handover.")
    document.add_paragraph("Feature flag rollout remains constrained until SRE approval is recorded.")
    document.add_paragraph("1. Confirm deployment window with SRE.", style="List Number")
    document.add_paragraph("2. Confirm rollback owner is on-call.", style="List Number")
    document.add_heading("Approval Matrix", level=2)
    table = document.add_table(rows=4, cols=3)
    table.cell(0, 0).text = "Check"
    table.cell(0, 1).text = "Owner"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "Security sign-off"
    table.cell(1, 1).text = "Security Lead"
    table.cell(1, 2).text = "PENDING"
    table.cell(2, 0).text = "Rollback readiness"
    table.cell(2, 1).text = "SRE"
    table.cell(2, 2).text = "READY"
    table.cell(3, 0).text = "Customer notification"
    table.cell(3, 1).text = "Product Owner"
    table.cell(3, 2).text = "APPROVED"
    document.add_heading("Known Limitations", level=2)
    document.add_paragraph("A rollback drill is documented, but the final go/no-go meeting must verify production readiness.")
    document.add_paragraph("Customer notification copy is ready and waiting for product owner sign-off.")
    document.add_heading("Appendix A: Rollback Contacts", level=2)
    document.add_paragraph("Primary on-call: sre-primary@example.com")
    document.add_paragraph("Secondary on-call: sre-secondary@example.com")
    document.save(path)


def _write_pdf(path: Path) -> None:
    try:
        import fitz
    except Exception as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("Для генерации PDF demo требуется зависимость PyMuPDF") from exc

    doc = fitz.open()
    page = doc.new_page()
    text = (
        "Audit Summary: Payments v2\n"
        "Security audit status: conditional pass.\n"
        "Critical vulnerabilities: none open.\n\n"
        "Control | Owner | Status | Evidence\n"
        "Security sign-off | Security Lead | PENDING | SEC-742\n"
        "Rollback drill | SRE | READY | OPS-113\n"
        "Customer notification | Product Owner | APPROVED | CRM-991\n\n"
        "Release Gate Form\n"
        "Approver: Security Lead; Decision: CONDITIONAL PASS; Ticket: SEC-742\n"
        "Owner: Release Manager; Due: 2026-04-30; Escalation: Required\n\n"
        "Release control notes:\n"
        "- Evidence pack must include security sign-off.\n"
        "- Operations dashboard and rollback runbook must be attached before final approval.\n"
    )
    page.insert_textbox(fitz.Rect(72, 72, 540, 760), text, fontsize=11)
    doc.save(path)
    doc.close()


def _write_scanned_pdf(path: Path) -> None:
    try:
        import fitz
    except Exception as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("Для генерации PDF demo требуется зависимость PyMuPDF") from exc

    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    doc.close()


def _write_scanned_pdf_sidecar(path: Path) -> None:
    text = (
        "Scanned Sign-off Record\n\n"
        "Security sign-off: APPROVED by release manager.\n"
        "Operations readiness: rollback runbook attached.\n"
        "Final note: scanned signature page recovered via OCR fallback.\n"
    )
    path.write_text(text, encoding="utf-8")


def _write_xlsx(path: Path) -> None:
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("Для генерации XLSX demo требуется зависимость openpyxl") from exc

    workbook = Workbook()
    approvals = workbook.active
    approvals.title = "Approval Tracker"
    approvals.append(["Check", "Owner", "Status", "Notes"])
    approvals.append(["Security sign-off", "Security Lead", "PENDING", "Waiting for final review"])
    approvals.append(["Rollback readiness", "SRE", "READY", "Rollback runbook attached"])
    approvals.append(["Customer notification", "Product Owner", "APPROVED", "Messaging ready"])

    risks = workbook.create_sheet("Risk Register")
    risks.append(["Risk", "Severity", "Owner", "Mitigation"])
    risks.append(["Payment timeout spike", "medium", "Ops", "Canary and alert tuning"])
    risks.append(["Approval lag", "high", "Release Manager", "Escalate pending approvers before freeze"])

    workbook.save(path)


def _write_pptx(path: Path) -> None:
    try:
        from pptx import Presentation
    except Exception as exc:  # pragma: no cover - depends on local environment
        raise RuntimeError("Для генерации PPTX demo требуется зависимость python-pptx") from exc

    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Release Briefing: Payments v2"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "Security sign-off: PENDING"
    body.add_paragraph().text = "Rollback readiness: READY"
    body.add_paragraph().text = "Customer notification: APPROVED"
    slide.notes_slide.notes_text_frame.text = (
        "Notes: финальный go/no-go требует закрыть security approval и подтвердить on-call roster."
    )

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Open Release Risks"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "Pending security governance confirmation"
    nested = body.add_paragraph()
    nested.text = "Mitigation: security lead review before freeze"
    nested.level = 1
    body.add_paragraph().text = "Canary rollback drill not signed-off"
    slide.notes_slide.notes_text_frame.text = (
        "Notes: escalation required if approvals remain pending at T-4h to release."
    )

    presentation.save(path)


if __name__ == "__main__":
    main()
