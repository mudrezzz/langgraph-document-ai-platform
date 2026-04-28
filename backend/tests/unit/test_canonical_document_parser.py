from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest

from domain_docs.parsing import CanonicalDocumentParser
from infra.docs.ocr_gateway import SidecarPdfOcrGateway
from scripts.build_binary_demo_documents import build_binary_demo_documents


def test_canonical_parser_extracts_markdown_structure(tmp_path: Path) -> None:
    source = tmp_path / "release_decision.md"
    source.write_text(
        "# Release\n\n## Decision\n\n- GO is blocked by pending approval\n- Security review is required\n",
        encoding="utf-8",
    )

    document = CanonicalDocumentParser().parse_path(source)

    assert document.file_type == "md"
    assert document.metadata_profile["document_type"] == "methodology"
    assert len(document.content_blocks) == 2
    assert document.content_blocks[0].block_type == "bullet"
    assert document.content_blocks[0].heading_path == ["Release", "Decision"]
    assert len(document.section_summaries) >= 1
    assert document.parser_quality.parser_family == "markdown"
    assert document.parser_quality.headings_total == 2
    assert document.parser_quality.lists_total == 2
    assert document.parser_quality.blocks_total == 2


def test_canonical_parser_extracts_json_values(tmp_path: Path) -> None:
    source = tmp_path / "approvals.json"
    source.write_text(
        json.dumps({"security": {"status": "PENDING"}, "approvers": ["SRE", "Risk"]}),
        encoding="utf-8",
    )

    document = CanonicalDocumentParser().parse_path(source)

    assert document.file_type == "json"
    assert document.metadata_profile["document_type"] == "governance"
    assert len(document.content_blocks) == 3
    assert any("security.status: PENDING" in block.text for block in document.content_blocks)
    assert document.quality_flags == []
    assert document.parser_quality.parser_family == "json"
    assert document.parser_quality.extraction_mode == "structured"
    assert document.parser_quality.issues == []


def test_canonical_parser_doc_id_is_stable_for_relative_and_absolute_paths(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_dir = tmp_path / "docs"
    source_dir.mkdir()
    source = source_dir / "release_decision.md"
    source.write_text("# Release\n\nGO is blocked by pending approval.", encoding="utf-8")

    parser = CanonicalDocumentParser()
    monkeypatch.chdir(tmp_path)

    relative_document = parser.parse_path(Path("docs") / "release_decision.md")
    absolute_document = parser.parse_path(source.resolve())

    assert relative_document.doc_id == absolute_document.doc_id


def test_canonical_parser_parses_release_demo_directory() -> None:
    repo_root = Path(__file__).resolve().parents[3]
    dataset_dir = repo_root / "backend" / "examples" / "cases" / "release_go_no_go_multifile_case" / "input"
    pytest.importorskip("openpyxl")
    pytest.importorskip("pptx")
    build_binary_demo_documents(output_dir=dataset_dir, overwrite=True)

    documents = CanonicalDocumentParser(pdf_ocr_gateway=SidecarPdfOcrGateway()).parse_dir(dataset_dir)

    assert len(documents) == 9
    assert {document.file_type for document in documents} == {"docx", "json", "md", "pdf", "txt", "xlsx", "pptx"}
    assert sum(len(document.content_blocks) for document in documents) >= 8
    assert any(document.parser_quality.parser_family == "docx" for document in documents)
    assert any(document.parser_quality.parser_family == "pdf" for document in documents)
    assert any(document.parser_quality.parser_family == "xlsx" for document in documents)
    assert any(document.parser_quality.parser_family == "pptx" for document in documents)


def test_canonical_parser_reports_docx_table_quality_flags(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    source = tmp_path / "release_notes.docx"
    document = Document()
    document.add_heading("Release Notes", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Check"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Security sign-off"
    table.cell(1, 1).text = "Pending"
    document.save(source)

    parsed = CanonicalDocumentParser().parse_path(source)

    assert "docx_tables_detected" in parsed.quality_flags
    assert parsed.parser_quality.tables_total == 1
    assert len(parsed.extracted_tables) == 1
    assert parsed.extracted_tables[0].title == "Release Notes"
    assert parsed.extracted_tables[0].columns == ["Check", "Status"]
    assert parsed.extracted_tables[0].rows[0]["Check"] == "Security sign-off"
    assert any(block.block_type == "table_row" for block in parsed.content_blocks)


def test_canonical_parser_keeps_docx_table_under_nearest_heading(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    source = tmp_path / "release_notes.docx"
    document = Document()
    document.add_heading("Release Notes", level=1)
    document.add_heading("Approval Matrix", level=2)
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Check"
    table.cell(0, 1).text = "Owner"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "Customer notification"
    table.cell(1, 1).text = "Product Owner"
    table.cell(1, 2).text = "APPROVED"
    document.add_heading("Appendix A: Contacts", level=2)
    document.add_paragraph("Primary: sre@example.com")
    document.save(source)

    parsed = CanonicalDocumentParser().parse_path(source)

    table_row = next(block for block in parsed.content_blocks if block.block_type == "table_row")
    assert table_row.heading_path == ["Approval Matrix"]
    assert parsed.extracted_tables[0].title == "Approval Matrix"


def test_canonical_parser_detects_docx_lists_and_appendix(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    source = tmp_path / "appendix_notes.docx"
    document = Document()
    document.add_heading("Main", level=1)
    document.add_paragraph("1. Review rollout", style="List Number")
    document.add_heading("Appendix A: Contacts", level=2)
    document.add_paragraph("Primary: sre@example.com")
    document.save(source)

    parsed = CanonicalDocumentParser().parse_path(source)

    assert parsed.parser_quality.lists_total >= 1
    assert "appendix_section_detected" in parsed.quality_flags
    assert any(block.block_type == "bullet" for block in parsed.content_blocks)


def test_canonical_parser_extracts_xlsx_sheet_tables(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    source = tmp_path / "release_tracker.xlsx"
    workbook = Workbook()
    tracker = workbook.active
    tracker.title = "Approval Tracker"
    tracker.append(["Check", "Owner", "Status"])
    tracker.append(["Customer notification", "Product Owner", "APPROVED"])
    risks = workbook.create_sheet("Risk Register")
    risks.append(["Risk", "Severity"])
    risks.append(["Approval lag", "high"])
    workbook.save(source)

    parsed = CanonicalDocumentParser().parse_path(source)

    assert parsed.file_type == "xlsx"
    assert parsed.parser_quality.parser_family == "xlsx"
    assert parsed.parser_quality.tables_total == 2
    assert parsed.parser_quality.headings_total == 2
    assert "xlsx_tables_detected" in parsed.quality_flags
    assert "xlsx_workbook_detected" in parsed.quality_flags
    assert len(parsed.extracted_tables) == 2
    assert parsed.extracted_tables[0].title == "Approval Tracker"
    assert parsed.extracted_tables[0].metadata["sheet_name"] == "Approval Tracker"
    first_row = next(block for block in parsed.content_blocks if block.block_type == "table_row")
    assert first_row.heading_path == ["Approval Tracker"]
    assert first_row.metadata["sheet_name"] == "Approval Tracker"
    assert first_row.metadata["table_id"] == "XLSX-T-1"


def test_canonical_parser_extracts_pptx_slides_and_notes(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from pptx import Presentation

    source = tmp_path / "release_briefing.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Release Briefing"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "Security approval pending"
    nested = body.add_paragraph()
    nested.text = "Mitigation: security lead review"
    nested.level = 1
    slide.notes_slide.notes_text_frame.text = "Speaker note: escalate at T-4h."
    presentation.save(source)

    parsed = CanonicalDocumentParser().parse_path(source)

    assert parsed.file_type == "pptx"
    assert parsed.parser_quality.parser_family == "pptx"
    assert parsed.parser_quality.extraction_mode == "slide_text"
    assert parsed.parser_quality.pages_total == 1
    assert parsed.parser_quality.headings_total == 1
    assert parsed.parser_quality.lists_total >= 1
    assert "pptx_slides_detected" in parsed.quality_flags
    assert "pptx_notes_detected" in parsed.quality_flags
    assert any(block.block_type == "slide_title" for block in parsed.content_blocks)
    assert any(block.block_type == "bullet" for block in parsed.content_blocks)
    assert any(block.block_type == "note" for block in parsed.content_blocks)
    assert any(block.metadata.get("slide_number") == 1 for block in parsed.content_blocks)


def test_canonical_parser_reports_missing_xlsx_dependency_when_needed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_import = __import__

    def fake_import(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == "openpyxl":
            raise ImportError("openpyxl unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", fake_import)
    source = tmp_path / "sample.xlsx"
    source.write_bytes(b"not-a-real-xlsx")

    with pytest.raises(RuntimeError, match="openpyxl"):
        CanonicalDocumentParser().parse_path(source)


def test_canonical_parser_reports_missing_pptx_dependency_when_needed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_import = __import__

    def fake_import(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == "pptx":
            raise ImportError("pptx unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", fake_import)
    source = tmp_path / "sample.pptx"
    source.write_bytes(b"not-a-real-pptx")

    with pytest.raises(RuntimeError, match="python-pptx"):
        CanonicalDocumentParser().parse_path(source)


def test_canonical_parser_reports_empty_pdf_as_ocr_candidate(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    import fitz

    source = tmp_path / "scan.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(source)
    pdf.close()

    parsed = CanonicalDocumentParser().parse_path(source)

    assert "pdf_no_extractable_text" in parsed.quality_flags
    assert "ocr_required" in parsed.quality_flags
    assert "ocr_not_available" in parsed.quality_flags
    assert parsed.parser_quality.parser_family == "pdf"
    assert parsed.parser_quality.pages_total == 1
    assert any(issue.code == "ocr_required" for issue in parsed.parser_quality.issues)


def test_canonical_parser_assigns_unique_block_ids_for_multi_page_pdf(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    import fitz

    source = tmp_path / "multipage.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "Page one readiness evidence.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Page two security approval pending.")
    pdf.save(source)
    pdf.close()

    parsed = CanonicalDocumentParser().parse_path(source)

    block_ids = [block.block_id for block in parsed.content_blocks]
    assert len(block_ids) >= 2
    assert len(block_ids) == len(set(block_ids))
    assert parsed.structure_tree.block_ids == block_ids
    assert {block.metadata.get("page_number") for block in parsed.content_blocks} == {1, 2}
    assert all(block.metadata.get("source_kind") == "page_block" for block in parsed.content_blocks)
    assert all(isinstance(block.metadata.get("reading_order_index"), int) for block in parsed.content_blocks)
    assert all(block.metadata.get("layout_kind") in {"paragraph", "table_like"} for block in parsed.content_blocks)


def test_canonical_parser_uses_sidecar_ocr_for_scanned_pdf(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    import fitz

    source = tmp_path / "scan.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(source)
    pdf.close()
    source.with_suffix(".pdf.ocr.txt").write_text(
        "Recovered sign-off via OCR\n\nOperations readiness approved.",
        encoding="utf-8",
    )

    parsed = CanonicalDocumentParser(pdf_ocr_gateway=SidecarPdfOcrGateway()).parse_path(source)

    assert "ocr_required" in parsed.quality_flags
    assert "ocr_applied" in parsed.quality_flags
    assert "ocr_provider:sidecar" in parsed.quality_flags
    assert parsed.content_blocks
    assert parsed.content_blocks[0].metadata["ocr_provider"] == "sidecar"
    assert parsed.content_blocks[0].metadata["source_kind"] == "page_block"
    assert parsed.content_blocks[0].metadata["layout_source"] == "ocr_text"
    assert parsed.content_blocks[0].metadata["reading_order_index"] == 1
    assert parsed.content_blocks[0].metadata["layout_kind"] in {"paragraph", "table_like"}
    assert any(issue.code == "ocr_applied" for issue in parsed.parser_quality.issues)


def test_canonical_parser_marks_pdf_table_like_layout_blocks(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    import fitz

    source = tmp_path / "table_like.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(
        fitz.Rect(72, 72, 520, 180),
        "Check | Owner | Status\nCustomer notification | Product Owner | APPROVED",
    )
    pdf.save(source)
    pdf.close()

    parsed = CanonicalDocumentParser().parse_path(source)

    assert "pdf_table_like_blocks_detected" in parsed.quality_flags
    assert "pdf_tables_extracted" in parsed.quality_flags
    assert parsed.parser_quality.tables_total == 1
    assert len(parsed.extracted_tables) == 1
    assert parsed.extracted_tables[0].columns == ["Check", "Owner", "Status"]
    assert parsed.extracted_tables[0].rows[0]["Owner"] == "Product Owner"
    assert any(block.block_type == "table_row" for block in parsed.content_blocks)
    assert any(block.metadata.get("layout_kind") == "table_like" for block in parsed.content_blocks)
    assert any(block.metadata.get("source_kind") == "table_row" for block in parsed.content_blocks)
    assert any(block.metadata.get("page_number") == 1 for block in parsed.content_blocks)
    assert any(issue.code == "pdf_table_like_blocks_detected" for issue in parsed.parser_quality.issues)
    assert any(issue.code == "pdf_tables_extracted" for issue in parsed.parser_quality.issues)


def test_canonical_parser_extracts_pdf_form_like_key_value_rows(tmp_path: Path) -> None:
    pytest.importorskip("fitz")
    import fitz

    source = tmp_path / "form_like.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(
        fitz.Rect(72, 72, 520, 220),
        (
            "Release Gate Form\n"
            "Approver: Security Lead; Decision: CONDITIONAL PASS; Ticket: SEC-742\n"
            "Owner: Release Manager; Due: 2026-04-30; Escalation: Required\n"
        ),
    )
    pdf.save(source)
    pdf.close()

    parsed = CanonicalDocumentParser().parse_path(source)

    assert "pdf_tables_extracted" in parsed.quality_flags
    assert "pdf_form_like_blocks_detected" in parsed.quality_flags
    assert parsed.extracted_tables
    assert any(table.metadata.get("pdf_table_kind") == "form_like" for table in parsed.extracted_tables)
    assert any(block.block_type == "table_row" for block in parsed.content_blocks)
    assert any(block.metadata.get("pdf_table_kind") == "form_like" for block in parsed.content_blocks)
    assert any(issue.code == "pdf_form_like_blocks_detected" for issue in parsed.parser_quality.issues)


def test_canonical_parser_reports_missing_docx_dependency_when_needed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_import = __import__

    def fake_import(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == "docx":
            raise ImportError("docx unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", fake_import)
    source = tmp_path / "sample.docx"
    source.write_bytes(b"not-a-real-docx")

    with pytest.raises(RuntimeError, match="python-docx"):
        CanonicalDocumentParser().parse_path(source)


def test_canonical_parser_reports_missing_pdf_dependency_when_needed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_import = __import__

    def fake_import(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == "fitz":
            raise ImportError("fitz unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", fake_import)
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"%PDF-1.4")

    with pytest.raises(RuntimeError, match="PyMuPDF"):
        CanonicalDocumentParser().parse_path(source)
