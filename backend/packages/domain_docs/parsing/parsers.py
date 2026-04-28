from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from domain_docs.parsing.ocr import PdfOcrGateway
from schemas.documents.contracts import (
    CanonicalContentBlock,
    CanonicalDocument,
    CanonicalTable,
    CanonicalSectionSummary,
    CanonicalStructureNode,
    ParserQualityIssue,
    ParserQualitySummary,
)

_SUPPORTED_EXTENSIONS = {".md", ".txt", ".json", ".docx", ".pdf", ".xlsx", ".pptx"}
_HEADING_RE = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*(?:[-*]|\d+\.)\s+(.+?)\s*$")
_WHITESPACE_RE = re.compile(r"\s+")


class CanonicalDocumentParser:
    """Parser MVP for text-like sources into canonical document contracts."""

    def __init__(self, *, pdf_ocr_gateway: PdfOcrGateway | None = None) -> None:
        self._pdf_ocr_gateway = pdf_ocr_gateway

    def supported_extensions(self) -> set[str]:
        return set(_SUPPORTED_EXTENSIONS)

    def parse_path(self, source_path: str | Path, *, version: str = "1") -> CanonicalDocument:
        path = Path(source_path)
        if not path.exists():
            raise FileNotFoundError(f"Документ не найден: {path}")
        if not path.is_file():
            raise ValueError(f"Путь документа должен быть файлом: {path}")

        extension = path.suffix.lower()
        if extension not in _SUPPORTED_EXTENSIONS:
            raise ValueError(f"Неподдерживаемый формат документа: {extension}")

        doc_id = _build_doc_id(path)
        quality_flags: list[str] = []
        metrics: dict[str, int | None] = {
            "pages_total": None,
            "headings_total": 0,
            "lists_total": 0,
            "tables_total": 0,
        }
        extraction_mode = "text"
        parser_family = "text"

        if extension == ".json":
            raw_text = path.read_text(encoding="utf-8")
            parser_family = "json"
            extraction_mode = "structured"
            blocks, structure = self._parse_json(raw_text, quality_flags=quality_flags, metrics=metrics)
        elif extension == ".xlsx":
            parser_family = "xlsx"
            extraction_mode = "structured"
            blocks, structure, tables = self._parse_xlsx(path, quality_flags=quality_flags, metrics=metrics)
            raw_text = "\n".join(block.text for block in blocks)
        elif extension == ".pptx":
            parser_family = "pptx"
            extraction_mode = "slide_text"
            blocks, structure = self._parse_pptx(path, quality_flags=quality_flags, metrics=metrics)
            raw_text = "\n".join(block.text for block in blocks)
            tables = []
        elif extension == ".docx":
            parser_family = "docx"
            extraction_mode = "structured"
            blocks, structure, tables = self._parse_docx(path, quality_flags=quality_flags, metrics=metrics)
            raw_text = "\n".join(block.text for block in blocks)
        elif extension == ".pdf":
            parser_family = "pdf"
            extraction_mode = "page_text"
            blocks, structure = self._parse_pdf(path, quality_flags=quality_flags, metrics=metrics)
            raw_text = "\n".join(block.text for block in blocks)
            tables = []
        else:
            raw_text = path.read_text(encoding="utf-8")
            parser_family = "markdown" if extension == ".md" else "text"
            extraction_mode = "markdown" if extension == ".md" else "text"
            blocks, structure = self._parse_text(
                raw_text,
                markdown=extension == ".md",
                quality_flags=quality_flags,
                metrics=metrics,
            )
            tables = []

        if extension == ".json":
            tables = []

        if not raw_text.strip():
            quality_flags.append("empty_document")
        if not blocks:
            quality_flags.append("no_content_blocks")

        summaries = _build_section_summaries(blocks, structure)
        unique_flags = _unique_keep_order(quality_flags)
        return CanonicalDocument(
            doc_id=doc_id,
            source_path=str(path),
            version=version,
            file_type=extension.lstrip("."),
            metadata_profile={
                "file_name": path.name,
                "file_size_bytes": path.stat().st_size,
                "project_id": "p1",
                "document_type": _infer_document_type(path),
                "tags": _build_tags(path),
                "parser": self.__class__.__name__,
                "supported_mime_family": "text",
            },
            structure_tree=structure,
            content_blocks=blocks,
            extracted_tables=tables,
            section_summaries=summaries,
            quality_flags=unique_flags,
            parser_quality=_build_parser_quality_summary(
                parser_family=parser_family,
                extraction_mode=extraction_mode,
                structure=structure,
                blocks=blocks,
                quality_flags=unique_flags,
                metrics=metrics,
            ),
        )

    def parse_dir(self, source_dir: str | Path, *, version: str = "1") -> list[CanonicalDocument]:
        root = Path(source_dir)
        if not root.exists():
            raise FileNotFoundError(f"Директория документов не найдена: {root}")
        if not root.is_dir():
            raise ValueError(f"Путь должен быть директорией: {root}")

        files = sorted(
            path
            for path in root.rglob("*")
            if path.is_file() and path.suffix.lower() in _SUPPORTED_EXTENSIONS and not _is_ocr_sidecar(path)
        )
        if not files:
            raise ValueError(f"В директории {root} нет поддерживаемых документов")
        return [self.parse_path(path, version=version) for path in files]

    def _parse_text(
        self,
        raw_text: str,
        *,
        markdown: bool,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode]:
        blocks: list[CanonicalContentBlock] = []
        root = CanonicalStructureNode(node_id="root", title="Document", level=0)
        heading_stack: list[tuple[int, str]] = []
        current_heading_path: list[str] = []
        current_node = root
        paragraph_buffer: list[str] = []

        def flush_paragraph() -> None:
            nonlocal paragraph_buffer
            text = _normalize_text(" ".join(paragraph_buffer))
            paragraph_buffer = []
            if not text:
                return
            block = _build_block(
                text=text,
                block_type="paragraph",
                index=len(blocks) + 1,
                heading_path=current_heading_path,
            )
            blocks.append(block)
            current_node.block_ids.append(block.block_id)

        for raw_line in raw_text.splitlines():
            heading_match = _HEADING_RE.match(raw_line) if markdown else None
            if heading_match:
                flush_paragraph()
                title = _normalize_text(heading_match.group(2))
                level = len(heading_match.group(1))
                heading_stack = [(item_level, item_title) for item_level, item_title in heading_stack if item_level < level]
                heading_stack.append((level, title))
                current_heading_path = [item_title for _, item_title in heading_stack]
                current_node = CanonicalStructureNode(
                    node_id=f"section-{len(root.children) + 1}",
                    title=title,
                    level=level,
                )
                root.children.append(current_node)
                metrics["headings_total"] = int(metrics.get("headings_total", 0) or 0) + 1
                continue

            bullet_match = _BULLET_RE.match(raw_line)
            if bullet_match:
                flush_paragraph()
                block = _build_block(
                    text=_normalize_text(bullet_match.group(1)),
                    block_type="bullet",
                    index=len(blocks) + 1,
                    heading_path=current_heading_path,
                )
                blocks.append(block)
                current_node.block_ids.append(block.block_id)
                metrics["lists_total"] = int(metrics.get("lists_total", 0) or 0) + 1
                continue

            if raw_line.strip():
                paragraph_buffer.append(raw_line.strip())
            else:
                flush_paragraph()

        flush_paragraph()

        if markdown and not root.children:
            quality_flags.append("no_structural_headings")
        if not markdown and blocks:
            root.block_ids = [block.block_id for block in blocks]
        return blocks, root

    def _parse_json(
        self,
        raw_text: str,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode]:
        try:
            payload = json.loads(raw_text)
        except json.JSONDecodeError as exc:
            quality_flags.append("json_parse_failed")
            raise ValueError(f"Некорректный JSON документ: {exc}") from exc

        values: list[tuple[str, str]] = []
        _walk_json(payload, path_prefix="", out=values)

        root = CanonicalStructureNode(node_id="root", title="JSON Document", level=0)
        blocks: list[CanonicalContentBlock] = []
        for index, (path, value) in enumerate(values, start=1):
            block = _build_block(
                text=f"{path}: {value}" if path else value,
                block_type="json_value",
                index=index,
                heading_path=[path] if path else [],
                metadata={"json_path": path},
            )
            blocks.append(block)
            root.block_ids.append(block.block_id)
        metrics["headings_total"] = len(root.children)
        return blocks, root

    def _parse_docx(
        self,
        path: Path,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode, list[CanonicalTable]]:
        try:
            from docx import Document
        except Exception as exc:  # pragma: no cover - depends on optional runtime package
            raise RuntimeError("Для DOCX parsing требуется зависимость python-docx") from exc

        document = Document(str(path))
        blocks: list[CanonicalContentBlock] = []
        tables: list[CanonicalTable] = []
        root = CanonicalStructureNode(node_id="root", title="DOCX Document", level=0)
        current_node = root
        current_heading_path: list[str] = []
        tables_detected = len(document.tables)
        if tables_detected:
            quality_flags.append("docx_tables_detected")
        metrics["tables_total"] = tables_detected

        table_index = 0
        for item_kind, item in _iter_docx_body_items(document):
            if item_kind == "paragraph":
                text = _normalize_text(item.text)
                if not text:
                    continue

                style_name = getattr(item.style, "name", "") or ""
                if style_name.lower().startswith("heading"):
                    current_heading_path = [text]
                    current_node = CanonicalStructureNode(
                        node_id=f"section-{len(root.children) + 1}",
                        title=text,
                        level=_parse_heading_level(style_name),
                    )
                    root.children.append(current_node)
                    metrics["headings_total"] = int(metrics.get("headings_total", 0) or 0) + 1
                    if _looks_like_appendix_heading(text):
                        quality_flags.append("appendix_section_detected")
                    continue

                block_type = "paragraph"
                if _looks_like_list_paragraph(style_name, text):
                    metrics["lists_total"] = int(metrics.get("lists_total", 0) or 0) + 1
                    block_type = "bullet"

                block = _build_block(
                    text=text,
                    block_type=block_type,
                    index=len(blocks) + 1,
                    heading_path=current_heading_path,
                    metadata={"style": style_name} if style_name else None,
                )
                blocks.append(block)
                current_node.block_ids.append(block.block_id)
                continue

            if item_kind == "table":
                table_index += 1
                parsed_table = _parse_docx_table(
                    item,
                    table_index=table_index,
                    current_heading_path=current_heading_path,
                    block_index_start=len(blocks) + 1,
                )
                tables.append(parsed_table.table)
                blocks.extend(parsed_table.blocks)
                current_node.block_ids.extend(block.block_id for block in parsed_table.blocks)

        if not root.children:
            quality_flags.append("no_structural_headings")
            root.block_ids = [block.block_id for block in blocks]
        return blocks, root, tables

    def _parse_xlsx(
        self,
        path: Path,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode, list[CanonicalTable]]:
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover - depends on optional runtime package
            raise RuntimeError("Для XLSX parsing требуется зависимость openpyxl") from exc

        workbook = load_workbook(filename=str(path), data_only=True)
        root = CanonicalStructureNode(node_id="root", title="XLSX Workbook", level=0)
        blocks: list[CanonicalContentBlock] = []
        tables: list[CanonicalTable] = []
        sheet_headings_total = 0
        table_index = 0

        for sheet in workbook.worksheets:
            rows = [
                [_normalize_text("" if cell is None else str(cell)) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell for cell in row)]
            if not rows:
                continue

            sheet_headings_total += 1
            sheet_node = CanonicalStructureNode(
                node_id=f"sheet-{sheet_headings_total}",
                title=sheet.title,
                level=1,
            )
            root.children.append(sheet_node)
            current_heading_path = [sheet.title]

            header = rows[0]
            data_rows = rows[1:] if len(rows) > 1 else []
            if not any(header):
                header = [f"column_{index + 1}" for index in range(len(rows[0]))]

            if data_rows:
                quality_flags.append("xlsx_tables_detected")
                table_index += 1
                metrics["tables_total"] = int(metrics.get("tables_total", 0) or 0) + 1
                parsed_table = _build_tabular_content(
                    header=header,
                    data_rows=data_rows,
                    table_id=f"XLSX-T-{table_index}",
                    title=sheet.title,
                    heading_path=current_heading_path,
                    block_index_start=len(blocks) + 1,
                    table_metadata={"sheet_name": sheet.title, "sheet_index": sheet_headings_total},
                    block_metadata={"sheet_name": sheet.title},
                )
                tables.append(parsed_table.table)
                blocks.extend(parsed_table.blocks)
                sheet_node.block_ids.extend(block.block_id for block in parsed_table.blocks)
            else:
                for row in rows:
                    text = _normalize_text("; ".join(cell for cell in row if cell))
                    if not text:
                        continue
                    block = _build_block(
                        text=text,
                        block_type="paragraph",
                        index=len(blocks) + 1,
                        heading_path=current_heading_path,
                        metadata={"sheet_name": sheet.title},
                    )
                    blocks.append(block)
                    sheet_node.block_ids.append(block.block_id)

        metrics["headings_total"] = sheet_headings_total
        if metrics["tables_total"]:
            quality_flags.append("xlsx_workbook_detected")
        if not root.children:
            quality_flags.append("empty_workbook")
        return blocks, root, tables

    def _parse_pdf(
        self,
        path: Path,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode]:
        try:
            import fitz
        except Exception as exc:  # pragma: no cover - depends on optional runtime package
            raise RuntimeError("Для PDF parsing требуется зависимость PyMuPDF") from exc

        root = CanonicalStructureNode(node_id="root", title="PDF Document", level=0)
        blocks: list[CanonicalContentBlock] = []
        with fitz.open(str(path)) as pdf_document:
            metrics["pages_total"] = len(pdf_document)
            for page_index, page in enumerate(pdf_document, start=1):
                page_blocks = self._extract_pdf_page_blocks(page=page, page_index=page_index)
                for block in page_blocks:
                    block.block_id = f"B-{len(blocks) + 1}"
                    blocks.append(block)
                    root.block_ids.append(block.block_id)

        if not blocks:
            quality_flags.append("pdf_no_extractable_text")
            quality_flags.append("ocr_required")
            ocr_blocks = self._run_pdf_ocr(path, quality_flags=quality_flags, metrics=metrics)
            if ocr_blocks:
                root.block_ids = [block.block_id for block in ocr_blocks]
                blocks = ocr_blocks
        if any(str(block.metadata.get("layout_kind", "")).strip() == "table_like" for block in blocks):
            quality_flags.append("pdf_table_like_blocks_detected")
        if blocks and len(blocks) < 2:
            quality_flags.append("low_text_density")
        return blocks, root

    def _extract_pdf_page_blocks(self, *, page: Any, page_index: int) -> list[CanonicalContentBlock]:
        extracted: list[CanonicalContentBlock] = []
        current_section_title: str | None = None
        raw_layout_blocks = page.get_text("blocks") or []
        layout_blocks = [raw_block for raw_block in raw_layout_blocks if len(raw_block) >= 5]
        layout_blocks.sort(key=lambda raw_block: (round(float(raw_block[1]), 2), round(float(raw_block[0]), 2)))

        for block_index, raw_block in enumerate(layout_blocks, start=1):
            x0, y0, x1, y1, raw_text = raw_block[:5]
            text = _normalize_text(str(raw_text or ""))
            if not text:
                continue

            heading_candidate = _infer_pdf_section_title(text)
            if heading_candidate is not None:
                current_section_title = heading_candidate

            heading_path = [current_section_title] if current_section_title else [f"Page {page_index}"]
            layout_kind = _infer_pdf_layout_kind(text)
            block = _build_block(
                text=text,
                block_type="paragraph",
                index=0,  # assigned below with stable global order
                heading_path=heading_path,
                metadata={
                    "source_kind": "page_block",
                    "page_number": page_index,
                    "reading_order_index": block_index,
                    "layout_kind": layout_kind,
                    "bbox": [round(float(x0), 2), round(float(y0), 2), round(float(x1), 2), round(float(y1), 2)],
                    "layout_source": "pdf_blocks",
                },
            )
            block.page_number = page_index
            extracted.append(block)

        # Fallback keeps previous behavior for PDFs where `blocks` is empty.
        if not extracted:
            page_text = page.get_text("text")
            for block_index, paragraph in enumerate(re.split(r"\n\s*\n+", page_text), start=1):
                text = _normalize_text(paragraph)
                if not text:
                    continue
                heading_candidate = _infer_pdf_section_title(text)
                if heading_candidate is not None:
                    current_section_title = heading_candidate
                heading_path = [current_section_title] if current_section_title else [f"Page {page_index}"]
                layout_kind = _infer_pdf_layout_kind(text)
                block = _build_block(
                    text=text,
                    block_type="paragraph",
                    index=0,  # assigned below with stable global order
                    heading_path=heading_path,
                    metadata={
                        "source_kind": "page_block",
                        "page_number": page_index,
                        "reading_order_index": block_index,
                        "layout_kind": layout_kind,
                        "layout_source": "pdf_text",
                    },
                )
                block.page_number = page_index
                extracted.append(block)
        return extracted

    def _run_pdf_ocr(
        self,
        path: Path,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> list[CanonicalContentBlock]:
        if self._pdf_ocr_gateway is None:
            quality_flags.append("ocr_not_available")
            return []

        ocr_result = self._pdf_ocr_gateway.extract_pdf_text(path)
        if ocr_result is None or not ocr_result.page_texts:
            quality_flags.append("ocr_text_not_recovered")
            return []

        quality_flags.append("ocr_applied")
        quality_flags.append(f"ocr_provider:{ocr_result.provider}")
        for warning in ocr_result.warnings:
            quality_flags.append(warning)

        blocks: list[CanonicalContentBlock] = []
        for page_index, page_text in enumerate(ocr_result.page_texts, start=1):
            for block_index, paragraph in enumerate(re.split(r"\n\s*\n+", page_text), start=1):
                text = _normalize_text(paragraph)
                if not text:
                    continue
                layout_kind = _infer_pdf_layout_kind(text)
                block = _build_block(
                    text=text,
                    block_type="paragraph",
                    index=len(blocks) + 1,
                    heading_path=[f"Page {page_index}"],
                    metadata={
                        "source_kind": "page_block",
                        "page_number": page_index,
                        "reading_order_index": block_index,
                        "layout_kind": layout_kind,
                        "ocr_provider": ocr_result.provider,
                        "extraction_mode": "ocr",
                        "layout_source": "ocr_text",
                    },
                )
                block.page_number = page_index
                blocks.append(block)

        metrics["pages_total"] = max(int(metrics.get("pages_total") or 0), len(ocr_result.page_texts))
        return blocks

    def _parse_pptx(
        self,
        path: Path,
        *,
        quality_flags: list[str],
        metrics: dict[str, int | None],
    ) -> tuple[list[CanonicalContentBlock], CanonicalStructureNode]:
        try:
            from pptx import Presentation
        except Exception as exc:  # pragma: no cover - depends on optional runtime package
            raise RuntimeError("Для PPTX parsing требуется зависимость python-pptx") from exc

        presentation = Presentation(str(path))
        root = CanonicalStructureNode(node_id="root", title="PPTX Presentation", level=0)
        blocks: list[CanonicalContentBlock] = []
        slides_total = len(presentation.slides)
        metrics["pages_total"] = slides_total
        if slides_total:
            quality_flags.append("pptx_slides_detected")

        notes_detected = False
        for slide_index, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            slide_title = (
                _normalize_text(title_shape.text)
                if title_shape is not None and _normalize_text(title_shape.text)
                else f"Slide {slide_index}"
            )
            slide_node = CanonicalStructureNode(
                node_id=f"slide-{slide_index}",
                title=slide_title,
                level=1,
            )
            root.children.append(slide_node)
            metrics["headings_total"] = int(metrics.get("headings_total", 0) or 0) + 1
            heading_path = [slide_title]

            if title_shape is not None and _normalize_text(title_shape.text):
                title_block = _build_block(
                    text=slide_title,
                    block_type="slide_title",
                    index=len(blocks) + 1,
                    heading_path=heading_path,
                    metadata={"slide_number": slide_index},
                )
                blocks.append(title_block)
                slide_node.block_ids.append(title_block.block_id)

            for shape_index, shape in enumerate(slide.shapes, start=1):
                if title_shape is not None and shape == title_shape:
                    continue
                if not getattr(shape, "has_text_frame", False):
                    continue
                text_frame = getattr(shape, "text_frame", None)
                if text_frame is None:
                    continue
                for paragraph in text_frame.paragraphs:
                    text = _normalize_text(paragraph.text or "")
                    if not text:
                        continue
                    level = int(getattr(paragraph, "level", 0) or 0)
                    block_type = "bullet" if level > 0 else "paragraph"
                    if block_type == "bullet":
                        metrics["lists_total"] = int(metrics.get("lists_total", 0) or 0) + 1
                    block = _build_block(
                        text=text,
                        block_type=block_type,
                        index=len(blocks) + 1,
                        heading_path=heading_path,
                        metadata={
                            "slide_number": slide_index,
                            "shape_index": shape_index,
                            "paragraph_level": level,
                        },
                    )
                    blocks.append(block)
                    slide_node.block_ids.append(block.block_id)

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes_text = _normalize_text(slide.notes_slide.notes_text_frame.text or "")
                if notes_text:
                    notes_detected = True
                    note_block = _build_block(
                        text=notes_text,
                        block_type="note",
                        index=len(blocks) + 1,
                        heading_path=heading_path,
                        metadata={"slide_number": slide_index},
                    )
                    blocks.append(note_block)
                    slide_node.block_ids.append(note_block.block_id)

        if notes_detected:
            quality_flags.append("pptx_notes_detected")
        if not root.children:
            quality_flags.append("empty_presentation")
        return blocks, root


def _build_parser_quality_summary(
    *,
    parser_family: str,
    extraction_mode: str,
    structure: CanonicalStructureNode,
    blocks: list[CanonicalContentBlock],
    quality_flags: list[str],
    metrics: dict[str, int | None],
) -> ParserQualitySummary:
    sections_total = len(structure.children)
    flags = list(quality_flags)
    issues = [_build_quality_issue(flag, metrics=metrics, blocks=blocks, sections_total=sections_total) for flag in flags]
    return ParserQualitySummary(
        parser_family=parser_family,
        extraction_mode=extraction_mode,
        pages_total=metrics.get("pages_total"),
        blocks_total=len(blocks),
        sections_total=sections_total,
        headings_total=int(metrics.get("headings_total", 0) or 0),
        lists_total=int(metrics.get("lists_total", 0) or 0),
        tables_total=int(metrics.get("tables_total", 0) or 0),
        issues=issues,
        flags=flags,
    )


def _build_quality_issue(
    flag: str,
    *,
    metrics: dict[str, int | None],
    blocks: list[CanonicalContentBlock],
    sections_total: int,
) -> ParserQualityIssue:
    if flag == "empty_document":
        return ParserQualityIssue(code=flag, severity="blocking", message="Document text is empty after parser read")
    if flag == "no_content_blocks":
        return ParserQualityIssue(code=flag, severity="blocking", message="Parser did not produce content blocks")
    if flag == "pdf_no_extractable_text":
        return ParserQualityIssue(
            code=flag,
            severity="blocking",
            message="PDF parser found no extractable text blocks",
            metadata={"pages_total": metrics.get("pages_total")},
        )
    if flag == "ocr_required":
        return ParserQualityIssue(
            code=flag,
            severity="warning",
            message="PDF likely requires OCR before reliable indexing",
            metadata={"pages_total": metrics.get("pages_total")},
        )
    if flag == "ocr_not_available":
        return ParserQualityIssue(
            code=flag,
            severity="warning",
            message="OCR fallback is not configured for this runtime",
        )
    if flag == "ocr_text_not_recovered":
        return ParserQualityIssue(
            code=flag,
            severity="blocking",
            message="OCR fallback did not recover text from scanned PDF",
        )
    if flag == "ocr_applied":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="OCR fallback recovered text for scanned PDF",
        )
    if flag.startswith("ocr_provider:"):
        provider = flag.split(":", 1)[1] or "unknown"
        return ParserQualityIssue(
            code="ocr_provider",
            severity="info",
            message=f"OCR provider used: {provider}",
            metadata={"provider": provider},
        )
    if flag == "low_text_density":
        return ParserQualityIssue(
            code=flag,
            severity="warning",
            message="Parser extracted very few text blocks for the document",
            metadata={"blocks_total": len(blocks), "pages_total": metrics.get("pages_total")},
        )
    if flag == "pdf_table_like_blocks_detected":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="PDF parser detected table-like logical blocks",
            metadata={"blocks_total": len(blocks), "pages_total": metrics.get("pages_total")},
        )
    if flag == "no_structural_headings":
        return ParserQualityIssue(
            code=flag,
            severity="warning",
            message="Parser did not detect structural headings",
            metadata={"sections_total": sections_total},
        )
    if flag == "docx_tables_detected":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="DOCX contains tables that were detected during parsing",
            metadata={"tables_total": metrics.get("tables_total", 0)},
        )
    if flag == "appendix_section_detected":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="Appendix-like section detected in DOCX structure",
        )
    if flag == "pptx_slides_detected":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="PPTX slides were detected and parsed",
            metadata={"slides_total": metrics.get("pages_total", 0)},
        )
    if flag == "pptx_notes_detected":
        return ParserQualityIssue(
            code=flag,
            severity="info",
            message="Speaker notes were detected in PPTX slides",
        )
    if flag == "empty_presentation":
        return ParserQualityIssue(
            code=flag,
            severity="blocking",
            message="Presentation contains no parseable slides",
        )
    if flag == "table_extraction_not_implemented":
        return ParserQualityIssue(
            code=flag,
            severity="warning",
            message="Tabular content was detected but table extraction is not implemented yet",
            metadata={"tables_total": metrics.get("tables_total", 0)},
        )
    return ParserQualityIssue(code=flag, severity="warning", message=f"Parser quality flag: {flag}")


def _looks_like_list_paragraph(style_name: str, text: str) -> bool:
    lowered_style = style_name.lower()
    return lowered_style.startswith("list") or bool(_BULLET_RE.match(text))


def _looks_like_appendix_heading(text: str) -> bool:
    lowered = text.strip().lower()
    return lowered.startswith(("appendix", "annex", "приложение"))


class _ParsedDocxTable:
    def __init__(self, *, table: CanonicalTable, blocks: list[CanonicalContentBlock]) -> None:
        self.table = table
        self.blocks = blocks


def _parse_docx_table(
    table: Any,
    *,
    table_index: int,
    current_heading_path: list[str],
    block_index_start: int,
) -> _ParsedDocxTable:
    rows = [[_normalize_text(cell.text) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return _ParsedDocxTable(
            table=CanonicalTable(table_id=f"T-{table_index}", title=f"Table {table_index}"),
            blocks=[],
        )

    header = rows[0]
    data_rows = rows[1:] if len(rows) > 1 else []
    return _build_tabular_content(
        header=header,
        data_rows=data_rows,
        table_id=f"T-{table_index}",
        title=current_heading_path[-1] if current_heading_path else f"Table {table_index}",
        heading_path=current_heading_path,
        block_index_start=block_index_start,
        table_metadata={"heading_path": list(current_heading_path), "table_index": table_index},
        block_metadata={},
    )


def _build_tabular_content(
    *,
    header: list[str],
    data_rows: list[list[str]],
    table_id: str,
    title: str,
    heading_path: list[str],
    block_index_start: int,
    table_metadata: dict[str, Any],
    block_metadata: dict[str, Any],
) -> _ParsedDocxTable:
    if not any(header):
        header = [f"column_{index + 1}" for index in range(len(header))]

    normalized_columns = [column or f"column_{index + 1}" for index, column in enumerate(header)]
    normalized_rows: list[dict[str, Any]] = []
    blocks: list[CanonicalContentBlock] = []
    for row_index, row in enumerate(data_rows, start=1):
        padded = list(row) + [""] * max(0, len(normalized_columns) - len(row))
        item = {normalized_columns[index]: padded[index] for index in range(len(normalized_columns))}
        normalized_rows.append(item)
        pairs = [f"{key}: {value}" for key, value in item.items() if value]
        row_text = _normalize_text("; ".join(pairs))
        if row_text:
            blocks.append(
                _build_block(
                    text=row_text,
                    block_type="table_row",
                    index=block_index_start + len(blocks),
                    heading_path=heading_path,
                    metadata={"table_id": table_id, "row_index": row_index, **block_metadata},
                )
            )

    return _ParsedDocxTable(
        table=CanonicalTable(
            table_id=table_id,
            title=title,
            columns=normalized_columns,
            rows=normalized_rows,
            metadata=dict(table_metadata),
        ),
        blocks=blocks,
    )


def _build_block(
    *,
    text: str,
    block_type: str,
    index: int,
    heading_path: list[str],
    metadata: dict[str, Any] | None = None,
) -> CanonicalContentBlock:
    return CanonicalContentBlock(
        block_id=f"B-{index}",
        block_type=block_type,
        text=text,
        heading_path=list(heading_path),
        metadata=metadata or {},
    )


def _walk_json(value: Any, *, path_prefix: str, out: list[tuple[str, str]]) -> None:
    if isinstance(value, dict):
        for key, nested in value.items():
            nested_prefix = f"{path_prefix}.{key}" if path_prefix else str(key)
            _walk_json(nested, path_prefix=nested_prefix, out=out)
        return
    if isinstance(value, list):
        for index, nested in enumerate(value):
            _walk_json(nested, path_prefix=f"{path_prefix}[{index}]", out=out)
        return
    if isinstance(value, (str, int, float, bool)):
        normalized = _normalize_text(str(value))
        if normalized:
            out.append((path_prefix, normalized))


def _build_section_summaries(
    blocks: list[CanonicalContentBlock],
    structure: CanonicalStructureNode,
) -> list[CanonicalSectionSummary]:
    if not blocks:
        return []

    summaries: list[CanonicalSectionSummary] = []
    nodes = structure.children or [structure]
    block_by_id = {block.block_id: block for block in blocks}

    for node in nodes:
        node_blocks = [block_by_id[block_id] for block_id in node.block_ids if block_id in block_by_id]
        if not node_blocks and node is structure:
            node_blocks = blocks[:3]
        if not node_blocks:
            continue
        summary_text = _normalize_text(" ".join(block.text for block in node_blocks))[:320]
        summaries.append(
            CanonicalSectionSummary(
                section_id=node.node_id,
                title=node.title,
                summary=summary_text,
                source_block_ids=[block.block_id for block in node_blocks],
            )
        )
    return summaries


def _build_doc_id(path: Path) -> str:
    stem = path.stem
    stable_name = path.name
    fingerprint = sum((index + 1) * ord(char) for index, char in enumerate(stable_name))
    prefix = re.sub(r"[^A-Za-z0-9]+", "", stem.upper())[:8] or "DOC"
    return f"{prefix}-{fingerprint % 10000:04d}"


def _parse_heading_level(style_name: str) -> int:
    match = re.search(r"(\d+)", style_name)
    if not match:
        return 1
    return int(match.group(1))


def _infer_document_type(path: Path) -> str:
    lowered = path.stem.lower()
    if any(token in lowered for token in ("briefing", "deck", "slide", "presentation")):
        return "governance"
    if any(token in lowered for token in ("tracker", "register", "matrix", "spreadsheet")):
        return "governance"
    if any(token in lowered for token in ("security", "sec", "vuln", "pentest")):
        return "security"
    if any(token in lowered for token in ("ops", "rollback", "monitor", "sre")):
        return "operations"
    if any(token in lowered for token in ("approval", "governance", "policy")):
        return "governance"
    if any(token in lowered for token in ("decision", "method", "strategy", "scope")):
        return "methodology"
    return "requirements"


def _build_tags(path: Path) -> list[str]:
    tags = []
    for token in re.split(r"[^a-z0-9]+", path.stem.lower()):
        normalized = token.strip()
        if normalized and normalized not in tags:
            tags.append(normalized)
    return tags


def _normalize_text(value: str) -> str:
    return _WHITESPACE_RE.sub(" ", value).strip()


def _infer_pdf_section_title(text: str) -> str | None:
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    if not lines:
        return None
    candidate = lines[0]
    normalized = _normalize_text(candidate)
    if len(normalized) < 3 or len(normalized) > 90:
        return None
    if normalized.endswith(":"):
        return normalized.rstrip(":")
    alpha = [ch for ch in normalized if ch.isalpha()]
    if alpha and sum(1 for ch in alpha if ch.isupper()) / float(len(alpha)) >= 0.7:
        return normalized.title()
    if re.match(r"^(section|chapter|part)\s+\d+", normalized, flags=re.IGNORECASE):
        return normalized
    return None


def _infer_pdf_layout_kind(text: str) -> str:
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    if "|" in text and text.count("|") >= 2:
        return "table_like"
    if re.search(r"\b\w+\s*:\s*[^;]+;\s*\w+\s*:\s*[^;]+", text):
        return "table_like"
    column_like_lines = sum(1 for line in lines if re.search(r"\S\s{2,}\S", line))
    if column_like_lines >= 2:
        return "table_like"
    return "paragraph"


def _iter_docx_body_items(document: Any) -> list[tuple[str, Any]]:
    try:
        from docx.document import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception as exc:  # pragma: no cover - depends on optional runtime package
        raise RuntimeError("Для DOCX parsing требуется зависимость python-docx") from exc

    items: list[tuple[str, Any]] = []
    parent = document if isinstance(document, DocxDocument) else document._body  # noqa: SLF001
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            items.append(("paragraph", Paragraph(child, parent)))
        elif isinstance(child, CT_Tbl):
            items.append(("table", Table(child, parent)))
    return items


def _is_ocr_sidecar(path: Path) -> bool:
    return path.name.endswith(".pdf.ocr.txt")


def _unique_keep_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result
